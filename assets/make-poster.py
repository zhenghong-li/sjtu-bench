"""Generate the SJTU-Bench promo poster (16:9) with python-pptx."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

W, H = 13.333, 7.5
BG = RGBColor(0x0B, 0x12, 0x20)
CARD = RGBColor(0x12, 0x1C, 0x2E)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
LIME = RGBColor(0xA3, 0xE6, 0x35)
ORANGE = RGBColor(0xFB, 0x92, 0x3C)
FONT = "Microsoft YaHei"


def set_font(run, size, bold=False, color=WHITE, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font)


def textbox(slide, left, top, width, height, lines, size=16, color=WHITE,
            bold=False, align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if isinstance(lines, str):
        lines = [lines]
    for index, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = spacing
        if isinstance(line, tuple):
            text, opts = line
        else:
            text, opts = line, {}
        run = paragraph.add_run()
        run.text = text
        set_font(run, opts.get("size", size), opts.get("bold", bold),
                 opts.get("color", color))
    return box


def rect(slide, left, top, width, height, fill, radius=True, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if line is None else None
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def circle(slide, left, top, size, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(size), Inches(size))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BG

# decorative shapes
circle(slide, 11.2, -1.4, 4.2, RGBColor(0x0F, 0x2A, 0x3A))
circle(slide, 12.0, 4.6, 2.6, RGBColor(0x1A, 0x2A, 0x12))
circle(slide, 10.4, 5.6, 1.4, RGBColor(0x2A, 0x1A, 0x10))
rect(slide, 0.0, 0.0, 0.16, H, CYAN, radius=False)
rect(slide, 0.7, 0.55, 3.3, 0.42, RGBColor(0x16, 0x2A, 0x3C))
textbox(slide, 0.82, 0.62, 3.1, 0.3, "SJTU · CAMPUS AI BENCH", 11, CYAN, True)

# title (single line, two runs)
title = slide.shapes.add_textbox(Inches(0.7), Inches(1.08), Inches(9.4), Inches(0.9))
title_tf = title.text_frame
title_tf.word_wrap = False
title_tf.margin_left = 0
title_tf.margin_right = 0
title_tf.margin_top = 0
title_tf.margin_bottom = 0
title_para = title_tf.paragraphs[0]
for text, color in [("SJTU-", WHITE), ("BENCH", CYAN)]:
    run = title_para.add_run()
    run.text = text
    set_font(run, 54, True, color)
textbox(slide, 0.72, 2.06, 8.4, 0.62, "15 个任务，测测你的校园 AI", 28, WHITE, True)
textbox(slide, 0.72, 2.74, 8.3, 0.5, "登一次 jAccount，看 AI 能不能真的帮你办事", 15, MUTED)

# left bullet list
bullets = [
    ("L1", "查课表 / 通知 / 水源 / 交我办", LIME),
    ("L2", "跨源汇总：DDL、公众号、邮箱、活动", CYAN),
    ("L3", "多步操作：政务流程 / 云盘 / 报销（只预填）", ORANGE),
    ("L4", "每日简报 / 工作流设计 / 你的 Agent 方案", RGBColor(0xC4, 0xB5, 0xFD)),
]
y = 3.52
for tag, desc, color in bullets:
    rect(slide, 0.72, y, 0.62, 0.36, color)
    textbox(slide, 0.72, y + 0.025, 0.62, 0.3, tag, 12, BG, True, PP_ALIGN.CENTER)
    textbox(slide, 1.48, y - 0.01, 6.7, 0.42, desc, 13, WHITE)
    y += 0.55

# right card: trigger
rect(slide, 8.45, 1.55, 4.25, 4.85, CARD)
rect(slide, 8.75, 1.85, 3.65, 0.52, CYAN)
textbox(slide, 8.75, 1.94, 3.65, 0.32, "对 agent 说这句话", 12, BG, True, PP_ALIGN.CENTER)
textbox(slide, 8.7, 2.62, 3.75, 1.0, "“我要试一试\nsjtubench”", 20, WHITE, True, PP_ALIGN.CENTER, 1.1)
textbox(slide, 8.72, 3.72, 3.7, 0.4, "↓ 自动跑完 15 个任务", 12, MUTED, False, PP_ALIGN.CENTER)

steps = ["1. 装好 sjtu-mcp + DSH", "2. 登录一次 jAccount", "3. 跑完可选上榜"]
y2 = 4.25
for step in steps:
    textbox(slide, 8.85, y2, 3.5, 0.32, step, 11.5, MUTED)
    y2 += 0.38

textbox(slide, 8.72, 5.75, 3.7, 0.4, "LLM 评分 · 匿名排行 · 隐私不出本机", 10.5, CYAN, False, PP_ALIGN.CENTER)

# footer
rect(slide, 0.0, 6.72, W, 0.78, RGBColor(0x0F, 0x18, 0x28), radius=False)
textbox(slide, 0.7, 6.88, 7.5, 0.4, "github.com/zhenghong-li/sjtu-bench", 14, WHITE, True)
textbox(slide, 8.6, 6.9, 4.0, 0.35, "娱乐向 · 非官方 · MIT", 11, MUTED, False, PP_ALIGN.RIGHT)

prs.save("assets/sjtu-bench-poster-16x9.pptx")
print("saved assets/sjtu-bench-poster-16x9.pptx")
