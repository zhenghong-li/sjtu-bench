"""Generate the SJTU-Bench 3:4 promo poster (小红书/朋友圈)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

W, H = 7.5, 10.0
BG = RGBColor(0x0B, 0x12, 0x20)
CARD = RGBColor(0x12, 0x1C, 0x2E)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
LIME = RGBColor(0xA3, 0xE6, 0x35)
ORANGE = RGBColor(0xFB, 0x92, 0x3C)
PURPLE = RGBColor(0xC4, 0xB5, 0xFD)
FONT = "Microsoft YaHei"


def set_font(run, size, bold=False, color=WHITE):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def textbox(slide, left, top, width, height, lines, size=14, color=WHITE,
            bold=False, align=PP_ALIGN.LEFT, spacing=1.0):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if isinstance(lines, str):
        lines = [lines]
    for index, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = spacing
        run = paragraph.add_run()
        if isinstance(line, tuple):
            text, opts = line
        else:
            text, opts = line, {}
        run.text = text
        set_font(run, opts.get("size", size), opts.get("bold", bold), opts.get("color", color))
    return box


def rect(slide, left, top, width, height, fill, radius=True):
    type_ = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(type_, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.1
    return shape


def circle(slide, left, top, size, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(size), Inches(size))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BG

circle(slide, 5.6, -1.2, 3.8, RGBColor(0x0F, 0x2A, 0x3A))
circle(slide, 6.2, 6.3, 2.0, RGBColor(0x1A, 0x2A, 0x12))
circle(slide, 5.7, 7.8, 1.4, RGBColor(0x2A, 0x1A, 0x10))
rect(slide, 0.0, 0.0, W, 0.12, CYAN, radius=False)

# badge
rect(slide, 0.55, 0.5, 2.9, 0.38, RGBColor(0x16, 0x2A, 0x3C))
textbox(slide, 0.68, 0.56, 2.7, 0.28, "SJTU · CAMPUS AI BENCH", 10, CYAN, True)

# title
title = slide.shapes.add_textbox(Inches(0.55), Inches(1.05), Inches(6.4), Inches(1.9))
tf = title.text_frame
tf.word_wrap = False
tf.margin_left = 0
tf.margin_right = 0
tf.margin_top = 0
tf.margin_bottom = 0
p1 = tf.paragraphs[0]
r1 = p1.add_run()
r1.text = "SJTU-"
set_font(r1, 54, True, WHITE)
p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = "BENCH"
set_font(r2, 54, True, CYAN)

# subtitle
sub = slide.shapes.add_textbox(Inches(0.58), Inches(2.95), Inches(6.3), Inches(0.55))
stf = sub.text_frame
stf.word_wrap = False
stf.margin_left = 0
stf.margin_right = 0
stf.margin_top = 0
stf.margin_bottom = 0
sp = stf.paragraphs[0]
for text, color in [("AI ", CYAN), ("帮你掌握校园的一切", WHITE)]:
    run = sp.add_run()
    run.text = text
    set_font(run, 23, True, color)
textbox(slide, 0.58, 3.55, 6.2, 0.3, "15 个真实校园任务 · L1–L4", 11.5, MUTED)

# four big task categories
levels = [
    ("L1", "查课表 / 通知 / 水源 / 交我办", LIME),
    ("L2", "DDL / 公众号 / 邮箱 / 活动", CYAN),
    ("L3", "政务流程 / 云盘 / 报销（只预填）", ORANGE),
    ("L4", "每日简报 / 工作流 / 你的 Agent 方案", PURPLE),
]
y = 3.98
for tag, desc, color in levels:
    rect(slide, 0.55, y, 6.4, 0.62, CARD)
    rect(slide, 0.55, y, 0.78, 0.62, color)
    textbox(slide, 0.55, y + 0.15, 0.78, 0.32, tag, 14, BG, True, PP_ALIGN.CENTER)
    textbox(slide, 1.55, y + 0.12, 5.2, 0.4, desc, 12.5, WHITE, True)
    y += 0.74

# trigger card (only requested elements)
rect(slide, 0.55, 6.95, 6.4, 2.17, CARD)
textbox(slide, 0.7, 7.3, 6.1, 0.5, "“我要试一试 sjtubench”", 18, WHITE, True, PP_ALIGN.CENTER)
textbox(slide, 0.7, 7.98, 6.1, 0.28, "登录一下 jAccount  ↓", 11, MUTED, False, PP_ALIGN.CENTER)
textbox(slide, 0.7, 8.3, 6.1, 0.4, "AI 能帮我省多少时间？", 15, WHITE, True, PP_ALIGN.CENTER)
textbox(slide, 0.7, 8.72, 6.1, 0.26, "一切隐私信息严格保存在本地", 9.5, CYAN, False, PP_ALIGN.CENTER)

# footer
rect(slide, 0.0, 9.3, W, 0.7, RGBColor(0x0F, 0x18, 0x28), radius=False)
textbox(slide, 0.55, 9.45, 4.6, 0.3, "github.com/zhenghong-li/sjtu-bench", 10.5, WHITE, True)
textbox(slide, 5.0, 9.46, 2.0, 0.28, "娱乐 · 非官方", 9, MUTED, False, PP_ALIGN.RIGHT)

prs.save("assets/sjtu-bench-poster-3x4.pptx")
print("saved assets/sjtu-bench-poster-3x4.pptx")
